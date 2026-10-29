"""表格规范化：区域网格 + 表头决策 → ``Table`` IR。

重复表头行不删除：IR 约定 ``rows[i]`` 对应绝对行 ``data_start_row + i``，
删行会让 footer、``full_grid()`` 与复核台改表头全部错位。重复行保留在
``rows`` 中，只把绝对行号记进 ``repeated_header_rows``，由
``Table.data_rows()`` 在数据视图里剔除。
"""
from __future__ import annotations

import re

from openpyxl.utils import get_column_letter

from ingest.config import Settings
from ingest.ir import Column, Table, Warning
from ingest.parsers.header_rules import HeaderCandidate, compose_header_names
from ingest.parsers.regions import Region
from ingest.parsers.xlsx_reader import SheetGrid

FOOTER_RE = re.compile(r"^(合\s*计|总\s*计|小\s*计|累计|备\s*注|注[:：]|说\s*明)")
UNIT_RE = re.compile(r"[（(]([^）)]{1,8})[）)]\s*$")
NUMBER_RE = re.compile(r"^[+-]?[\d,]+(\.\d+)?%?$")
DATE_RE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}")


def _is_number(text: str) -> bool:
    return bool(NUMBER_RE.match(text.replace("￥", "").replace("¥", "").strip()))


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _clean_name(text: str, index: int) -> tuple[str, str | None]:
    unit_match = UNIT_RE.search(text)
    unit = unit_match.group(1) if unit_match else None
    name = UNIT_RE.sub("", text).strip() or f"列{index}"
    return name, unit


def _infer_dtype(values: list[str]) -> str:
    values = [v for v in values if v.strip()]
    if not values:
        return "string"
    numeric = sum(1 for v in values if _is_number(v)) / len(values)
    if numeric >= 0.8:
        return "number"
    dates = sum(1 for v in values if DATE_RE.match(v)) / len(values)
    if dates >= 0.8:
        return "date"
    if all(v in {"是", "否", "true", "false", "True", "False"} for v in values):
        return "bool"
    return "string"


def _detect_footer_rows(sheet: SheetGrid, region: Region,
                        start_row: int) -> list[int]:
    footers: list[int] = []
    for r in range(start_row, region.end_row + 1):
        texts = [sheet.cell(r, c).text.strip()
                 for c in range(region.start_col, region.end_col + 1)]
        nonempty = [t for t in texts if t]
        first = next((t for t in texts[:3] if t), "")
        if FOOTER_RE.match(_norm(first)):
            footers.append(r)
            continue
        if len(nonempty) >= 2 and all(_is_number(t) for t in nonempty) \
                and any(not t for t in texts):
            footers.append(r)
    return footers


def _header_signature(header_values: list[list[str]]) -> list[set[str]]:
    """逐行表头的非空归一化单元格集合，按行顺序组成整块签名。

    重复表头是多行表头整体重复，只比对最后一行会把父级行留在数据视图里。
    """
    return [{_norm(v) for v in row if _norm(v)} for row in header_values]


def _repeated_header_rows(rows: list[list[str]], signature: list[set[str]],
                          data_start: int, *, excluded: set[int]) -> list[int]:
    """按完整表头块（逐行集合、顺序）匹配连续重复行；footer 优先，不再标重复。"""
    depth = len(signature)
    if depth == 0 or not any(signature):
        return []
    row_sets = [{_norm(v) for v in row if _norm(v)} for row in rows]
    marked: list[int] = []
    offset = 0
    while offset + depth <= len(rows):
        window = range(data_start + offset, data_start + offset + depth)
        if not any(r in excluded for r in window) and all(
                row_sets[offset + i] == signature[i] for i in range(depth)):
            marked.extend(window)
            offset += depth
        else:
            offset += 1
    return marked


def _columns_from_config(forced_columns: list[dict]) -> list[Column]:
    return [Column(
        key=cfg.get("key", f"c{i}"), name=cfg.get("name", f"列{i}"),
        dtype=cfg.get("dtype", "string"), unit=cfg.get("unit"),
        source_names=cfg.get("source_names", []),
    ) for i, cfg in enumerate(forced_columns, start=1)]


def build_table(sheet: SheetGrid, region: Region, header: HeaderCandidate | None, *,
                table_key: str, settings: Settings,
                forced_columns: list[dict] | None = None) -> tuple[Table, list[Warning]]:
    warnings: list[Warning] = []
    ncols = region.ncols
    header_rows = (
        [header.start_row + i for i in range(header.depth)] if header else []
    )
    header_values = [
        [sheet.cell(r, c).text for c in range(region.start_col, region.end_col + 1)]
        for r in header_rows
    ]
    # range 顶部、表头之前的行（标题/单位说明）：复核台前移表头时要能还原
    pre_rows = [
        [sheet.cell(r, c).text for c in range(region.start_col, region.end_col + 1)]
        for r in range(region.start_row, header.start_row)
    ] if header else []
    names = (compose_header_names(sheet, region, header.start_row, header.depth)
             if header else [f"列{i}" for i in range(1, ncols + 1)])

    data_start = max(header_rows) + 1 if header_rows else region.start_row
    raw_rows: list[list[str]] = [
        [sheet.cell(r, c).text for c in range(region.start_col, region.end_col + 1)]
        for r in range(data_start, region.end_row + 1)
    ]
    footer_rows = _detect_footer_rows(sheet, region, data_start)

    repeated_rows = _repeated_header_rows(
        raw_rows, _header_signature(header_values), data_start,
        excluded=set(footer_rows))
    if repeated_rows:
        warnings.append(Warning(
            code="repeated_header",
            detail=f"正文中重复表头行 {repeated_rows}，已从数据行中排除"))

    columns: list[Column] = []
    if forced_columns:
        columns = _columns_from_config(forced_columns)
    else:
        repeated = set(repeated_rows)
        for i, name in enumerate(names):
            clean, unit = _clean_name(name, i + 1)
            # 重复表头不是数据，不参与 dtype 推断，否则小表会被拖成 string
            values = [row[i] for offset, row in enumerate(raw_rows)
                      if data_start + offset not in repeated and i < len(row)]
            columns.append(Column(
                key=f"c{i + 1}", name=clean, dtype=_infer_dtype(values),
                unit=unit, source_names=[p for p in name.split("/") if p],
            ))
        if not columns:
            columns = [Column(key=f"c{i}", name=f"列{i}") for i in range(1, ncols + 1)]

    if len({c.name for c in columns}) != len(columns):
        seen: dict[str, int] = {}
        for i, col in enumerate(columns):
            count = seen.get(col.name, 0) + 1
            seen[col.name] = count
            if count > 1:
                columns[i] = col.model_copy(update={"name": f"{col.name}_{count}"})
        warnings.append(Warning(
            code="ambiguous_header", detail="列名重复，已加后缀区分"))

    if header is None:
        warnings.append(Warning(
            code="table_no_header", detail="未识别到表头，使用列1..列N"))
    if sheet.merges:
        warnings.append(Warning(
            code="merged_cells", detail=f"{len(sheet.merges)} 个合并区间"))

    confidence = header.confidence if header else 0.0
    source = header.source if header else "heuristic"
    if forced_columns and header:
        source = "template"
        confidence = 0.95
    if repeated_rows:
        confidence = min(confidence, 0.7)

    table = Table(
        table_key=table_key,
        sheet=sheet.name,
        range=(f"{get_column_letter(region.start_col)}{region.start_row}:"
               f"{get_column_letter(region.end_col)}{region.end_row}"),
        header_rows=header_rows,
        header_values=header_values,
        pre_rows=pre_rows,
        columns=columns,
        rows=raw_rows,
        footer_rows=footer_rows,
        repeated_header_rows=repeated_rows,
        merge_ranges=[m for m in sheet.merges],
        confidence=max(0.0, min(1.0, confidence)),
        source=source,
        flags=[w.code for w in warnings],
        candidates=[header.to_dict()] if header else [],
    )
    return table, warnings


def rebuild_table(table: Table, new_header_rows: list[int],
                  columns: list[dict], *, settings: Settings) -> tuple[Table, list[Warning]]:
    warnings: list[Warning] = []
    if not new_header_rows or new_header_rows != sorted(set(new_header_rows)):
        raise ValueError("new_header_rows 必须非空、递增且不重复")
    start = new_header_rows[0]
    depth = len(new_header_rows)
    # Table 的 rows[i] ↔ data_start_row + i 契约要求表头连续；非连续行号会让
    # data_start 与切出的数据行错位，因此必须显式拒绝（model_copy 不跑校验器）。
    if new_header_rows != list(range(start, start + depth)):
        raise ValueError("new_header_rows 必须是连续行号")
    if len(columns) != len(table.columns):
        raise ValueError(
            f"columns 数量（{len(columns)}）必须等于表格列数（{len(table.columns)}）")
    grid = table.full_grid()
    offset = start - table.range_start_row
    if offset < 0 or offset + depth > len(grid):
        raise ValueError("表头行超出表格范围")
    header_values = [list(row) for row in grid[offset:offset + depth]]
    pre_rows = [list(row) for row in grid[:offset]]
    rows = [list(row) for row in grid[offset + depth:]]
    data_start = start + depth
    footer_rows = [r for r in table.footer_rows if r >= data_start]
    # 旧标记可能已落进新表头范围，必须按新表头块在 rows 上重算
    repeated_rows = _repeated_header_rows(
        rows, _header_signature(header_values), data_start,
        excluded=set(footer_rows))
    if repeated_rows:
        warnings.append(Warning(
            code="repeated_header",
            detail=f"新表头下正文中仍有重复表头行 {repeated_rows}"))
    new_columns = _columns_from_config(columns)
    # 自动检测得到的三个表头状态码已不代表人工确认后的结果，先全部摘掉；
    # 重复表头按重算结果加回，其余代码（merged_cells/truncated 等）保留。
    flags = set(table.flags) - {"table_no_header", "ambiguous_header", "repeated_header"}
    if repeated_rows:
        flags.add("repeated_header")
    flags.add("human_confirmed")
    # 用 model_validate 重建以跑完整校验器（model_copy 会跳过 Table._check），
    # 保证 rows/header_rows/pre_rows 的绝对行号契约始终成立。
    rebuilt = Table.model_validate({
        **table.model_dump(),
        "header_rows": new_header_rows,
        "header_values": header_values,
        "pre_rows": pre_rows,
        "columns": new_columns,
        "rows": rows,
        "footer_rows": footer_rows,
        "repeated_header_rows": repeated_rows,
        "confidence": 1.0,
        "source": "human",
        "flags": sorted(flags),
    })
    return rebuilt, warnings
