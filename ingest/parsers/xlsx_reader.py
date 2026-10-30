from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from ingest.config import Settings, get_settings

# 每个工作簿最多记录多少条公式无缓存警告；超出的公式仍保留 formula 元数据，
# 只是不再刷屏（真实台账里动辄上千个 =SUM(...)）。
_MAX_FORMULA_WARNINGS = 20

_DECIMALS = re.compile(r"\.(0+)")


@dataclass(frozen=True)
class Cell:
    text: str = ""
    value: object = None
    is_numeric: bool = False
    bold: bool = False
    has_fill: bool = False
    hidden: bool = False
    formula: str | None = None


@dataclass
class SheetGrid:
    name: str
    hidden: bool
    rows: list[list[Cell]]
    merges: list[str]
    styles_available: bool = True

    @property
    def nrows(self) -> int:
        return len(self.rows)

    @property
    def ncols(self) -> int:
        return len(self.rows[0]) if self.rows else 0

    def cell(self, row: int, col: int) -> Cell:
        # 1-based 闭区间；越界返回空 Cell，调用方（区域切分/表头规则）无需防守
        if 1 <= row <= self.nrows and 1 <= col <= self.ncols:
            return self.rows[row - 1][col - 1]
        return Cell()


@dataclass
class WorkbookGrid:
    sheets: list[SheetGrid]
    warnings: list[dict]


def _decimals(fmt: str) -> int:
    match = _DECIMALS.search(fmt)
    return len(match.group(1)) if match else 0


def _display(cell) -> str:
    """尽量还原 Excel 里“看到的值”，而不是 Python 的 repr。

    只处理 P0 高频格式（百分比、千分位）；其他数字格式退回 str，
    避免自造格式化器把原文加工错。
    """
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        fmt = cell.number_format or ""
        if "%" in fmt:
            return f"{value * 100:.{_decimals(fmt)}f}%"
        if "#,##" in fmt:
            return f"{value:,.{_decimals(fmt)}f}"
        return str(value)
    return str(value)


def _is_numeric_value(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _hidden_columns(ws) -> set[int]:
    hidden: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if not dim.hidden:
            continue
        if dim.min and dim.max:
            # 分组列（如 A:C）只有一个 ColumnDimension，min/max 覆盖整段
            hidden.update(range(dim.min, dim.max + 1))
            continue
        try:
            hidden.add(column_index_from_string(str(key)))
        except ValueError:
            continue
    return hidden


def _hidden_rows(ws) -> set[int]:
    return {int(index) for index, dim in ws.row_dimensions.items() if dim.hidden}


def has_pivot_tables(path: Path) -> bool:
    """
    透视表只缓存计算结果，透视结构无法从 xlsx 复原。

    这里只看 zip 目录（`xl/pivotTables/`），不解析内容：
    openpyxl 的 data_only 模式拿到的就是缓存值，标记出来让复核人知道
    「表里的数是快照，不是可重算的透视」。
    """
    try:
        with zipfile.ZipFile(path) as zf:
            return any(name.startswith("xl/pivotTables/") for name in zf.namelist())
    except zipfile.BadZipFile:
        return False


def read_workbook(path: Path, settings: Settings | None = None) -> WorkbookGrid:
    s = settings or get_settings()
    # 两遍读取：data_only 拿缓存值（展示文本），默认模式拿公式原文。
    # 传文件对象而不是路径：内容寻址存储的原件没有扩展名，openpyxl 按扩展名
    # 判断格式会抛 InvalidFileException；文件对象则按 zip 头识别 xlsx。
    with Path(path).open("rb") as fh:
        wb_values = load_workbook(fh, data_only=True)
    with Path(path).open("rb") as fh:
        wb_formulas = load_workbook(fh, data_only=False)
    warnings: list[dict] = []
    if has_pivot_tables(path):
        warnings.append({"code": "pivot_values_only",
                         "detail": "工作簿含透视表，仅读取缓存值"})
    sheets: list[SheetGrid] = []
    remaining_cells = s.max_cells_per_workbook
    formula_warnings = 0

    for ws in wb_values.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        if max_row == 0 or max_col == 0:
            continue
        if remaining_cells <= 0:
            warnings.append({"code": "truncated",
                             "detail": "工作簿超过单元格上限，后续工作表已跳过"})
            break
        limit = min(s.max_cells_per_sheet, remaining_cells)
        if limit <= 0:
            warnings.append({"code": "truncated",
                             "detail": f"{ws.title} 超过单元格上限，已跳过"})
            break
        if max_row * max_col > limit:
            # 优先保列：行可以少读（Task 18 按行窗口切片），列一旦丢掉，
            # 表结构（Task 9 的 Table）就不可逆地缺字段，所以先截行。
            max_row = max(1, limit // max_col)
            if max_row * max_col > limit:
                # 单行都放不下时才截列（防御分支）
                max_col = max(1, limit // max_row)
            warnings.append({"code": "truncated",
                             "detail": f"{ws.title} 超过单元格上限，已截断"})
        remaining_cells -= max_row * max_col

        ws_formula = wb_formulas[ws.title]
        hidden_cols = _hidden_columns(ws)
        hidden_rows = _hidden_rows(ws)
        hidden_any = bool(hidden_cols or hidden_rows)
        rows: list[list[Cell]] = []
        for r in range(1, max_row + 1):
            row_cells: list[Cell] = []
            row_hidden = r in hidden_rows
            for c in range(1, max_col + 1):
                value_cell = ws.cell(r, c)
                formula_cell = ws_formula.cell(r, c)
                formula = (
                    formula_cell.value
                    if isinstance(formula_cell.value, str)
                    and formula_cell.value.startswith("=")
                    else None
                )
                if formula and value_cell.value is None \
                        and formula_warnings < _MAX_FORMULA_WARNINGS:
                    formula_warnings += 1
                    warnings.append({
                        "code": "formula_not_cached",
                        "detail": f"{ws.title}!{value_cell.coordinate}",
                    })
                row_cells.append(Cell(
                    text=_display(value_cell),
                    value=value_cell.value,
                    is_numeric=_is_numeric_value(value_cell.value),
                    bold=bool(value_cell.font.bold),
                    has_fill=bool(value_cell.fill.fill_type == "solid"),
                    hidden=row_hidden or c in hidden_cols,
                    formula=formula,
                ))
            rows.append(row_cells)

        # 合并区只有左上角有值；展开成完整网格，后续按绝对行号切表头才不丢列。
        # 合并区可能超出被截断的读取窗口，必须夹到窗口内，否则 IndexError。
        for merge in list(ws.merged_cells.ranges):
            if merge.min_row > max_row or merge.min_col > max_col:
                continue
            end_row = min(merge.max_row, max_row)
            end_col = min(merge.max_col, max_col)
            top = rows[merge.min_row - 1][merge.min_col - 1]
            for r in range(merge.min_row, end_row + 1):
                for c in range(merge.min_col, end_col + 1):
                    rows[r - 1][c - 1] = replace(
                        top, formula=None,
                        hidden=top.hidden or r in hidden_rows or c in hidden_cols)

        if hidden_any:
            warnings.append({"code": "hidden_rows",
                             "detail": f"{ws.title}：隐藏行列已标记"})

        sheets.append(SheetGrid(
            name=ws.title,
            hidden=ws.sheet_state != "visible",
            rows=rows,
            merges=[str(m) for m in ws.merged_cells.ranges],
            styles_available=True,
        ))

    return WorkbookGrid(sheets=sheets, warnings=warnings)
