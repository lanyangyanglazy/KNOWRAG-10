"""DOCX 解析编排：标题层级、段落、表格、文本框、修订与页眉页脚日期 → DocModel。

表格走与 XLSX 相同的表头链路（设计 §8.5）：模板 > 规则 > LLM。模板命中优先
于规则接受（人工确认过的版式比启发式分数更可信），LLM 每份文档最多 3 次且
只在 ``header_llm_enabled``（HEADER_LLM=always）时调用。Word 没有工作表名，
模板签名的 ``sheet_name`` 固定为 None。

``.doc`` 先经 LibreOffice 转成 docx 再复用同一套解析；转换失败抛
``ParseError``，由 worker 记为 failed。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from openpyxl.utils import get_column_letter

from ingest.config import Settings, get_settings
from ingest.convert import ConversionError, convert_document
from ingest.ir import Block, DocMeta, DocModel, ParserInfo, Warning
from ingest.parsers.base import ParseError
from ingest.parsers.header_llm import try_llm_header
from ingest.parsers.header_rules import (
    HeaderCandidate,
    compose_header_names,
    detect_header,
)
from ingest.parsers.regions import Region
from ingest.parsers.table_build import build_table
from ingest.parsers.templates import build_signature, match_template, template_columns
from ingest.parsers.xlsx_reader import Cell, SheetGrid
from ingest.sniff import sha256_file

_MAX_LLM_HEADER_CALLS = 3
# 与 XlsxParser 相同：规则候选只保留 top 3，纯数字表头甚至不产生候选，
# 模板兜底扫描限定在区域顶部 40 行 × 深度 1..3。
_MAX_TEMPLATE_SCAN_ROWS = 40
_MAX_TEMPLATE_DEPTH = 3

_HEADING_STYLE = re.compile(r"(?:Heading|标题)\s*(\d)", re.IGNORECASE)
_DATE_RE = re.compile(r"(20\d{2})[-年/](\d{1,2})[-月/](\d{1,2})")


def _iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, doc)


def _paragraph_text(paragraph: Paragraph) -> str:
    """段落正文 + VML 文本框文字（文本框在段落 XML 里是嵌套的 w:txbxContent）。"""
    parts = [paragraph.text.strip()]
    for textbox in paragraph._p.iter(qn("w:txbxContent")):
        texts = [t.text for t in textbox.iter(qn("w:t")) if t.text]
        joined = "".join(texts).strip()
        if joined and joined not in parts:
            parts.append(joined)
    return "\n".join(p for p in parts if p)


def _heading_level(paragraph: Paragraph, text: str) -> int | None:
    """标题层级：样式名优先，其次中文公文编号（一、/（一）/第X章）。"""
    match = _HEADING_STYLE.search(paragraph.style.name or "")
    if match:
        return min(6, int(match.group(1)))
    if re.match(r"^第[一二三四五六七八九十百]+[章节]", text):
        return 1
    if re.match(r"^[一二三四五六七八九十]+、", text):
        return 2
    if re.match(r"^（[一二三四五六七八九十]+）", text):
        return 3
    if re.match(r"^\d+(\.\d+)*[、.]", text):
        return 3
    return None


def _table_grid(table: DocxTable) -> SheetGrid:
    """Word 表格 → 与 XLSX 读取层同构的网格（1-based 闭区间，合并区已展开）。"""
    nrows = len(table.rows)
    ncols = len(table.columns)
    cells = [[table.cell(r, c) for c in range(ncols)] for r in range(nrows)]
    # 先缓存 _tc 引用再取 id：lxml 只保证代理存活期间同一节点对应同一代理，
    # 每次都用临时 cell 取 id 会让合并格算出不同 key，merge 区间就丢了。
    # python-docx 会把合并区展开到网格的每个位置（vMerge/gridSpan），
    # 因此同一 tc 覆盖的格子在这里会收敛成同一个 box。
    anchors = [[cell._tc for cell in row] for row in cells]
    rows = [[Cell(text=cell.text) for cell in row] for row in cells]
    bounds: dict[int, list[int]] = {}
    for r in range(nrows):
        for c in range(ncols):
            key = id(anchors[r][c])
            box = bounds.get(key)
            if box is None:
                bounds[key] = [r, c, r, c]
            else:
                box[2] = max(box[2], r)
                box[3] = max(box[3], c)
    merges = [
        f"{get_column_letter(c0 + 1)}{r0 + 1}:{get_column_letter(c1 + 1)}{r1 + 1}"
        for (r0, c0, r1, c1) in bounds.values() if (r0, c0) != (r1, c1)
    ]
    return SheetGrid(name="", hidden=False, rows=rows, merges=merges,
                     styles_available=False)


def _usable_templates(templates: list[dict], ncols: int) -> list[dict]:
    """只保留列数与区域宽度一致的模板；坏模板不能进入 build_table。"""
    return [t for t in templates if len(template_columns(t)) == ncols]


def _template_at(grid: SheetGrid, region: Region, start_row: int, depth: int,
                 templates: list[dict]) -> dict | None:
    candidate = HeaderCandidate(
        start_row=start_row, depth=depth, confidence=0.0,
        features={}, penalties=(), source="template")
    names = compose_header_names(grid, region, start_row, depth)
    signature = build_signature(None, region, candidate, names)
    return match_template(signature, templates)


def _header_footer_date(document) -> str | None:
    """从各节页眉/页脚里找第一个 ``YYYY-MM-DD`` 形态的日期。"""
    texts: list[str] = []
    for section in document.sections:
        for container in (section.header, section.footer):
            texts.extend(p.text for p in container.paragraphs)
    match = _DATE_RE.search("\n".join(texts))
    if not match:
        return None
    year, month, day = match.groups()
    return f"{year}-{int(month):02d}-{int(day):02d}"


class DocxParser:
    name = "docx"
    version = "0.1.0"

    def __init__(self, templates: list[dict] | None = None, llm=None,
                 settings: Settings | None = None):
        self.templates = templates or []
        self.llm = llm
        self.settings = settings or get_settings()

    def supports(self, sniff) -> bool:
        return sniff.kind in ("docx", "doc")

    def _resolve_header(self, grid: SheetGrid, region: Region,
                        decision) -> tuple[HeaderCandidate | None, list[dict] | None]:
        # 模板优先（设计 §8.5）：命中即用，不能等规则“拒绝”了才回头看模板，
        # 否则规则一旦接受别的行，模板就永远不生效。列数不符的模板直接跳过。
        templates = _usable_templates(self.templates, region.ncols)
        for cand in decision.candidates:
            template = _template_at(
                grid, region, cand.start_row, cand.depth, templates)
            if template:
                return (
                    HeaderCandidate(
                        start_row=cand.start_row, depth=cand.depth,
                        confidence=0.95, features=dict(cand.features),
                        penalties=(), source="template"),
                    template_columns(template),
                )

        # 候选为空（纯数字表头）或被 top 3 挤掉时，按区域顶部 40 行 ×
        # 深度 1..3 兜底扫描模板，先命中者胜。
        if templates:
            last_row = min(region.end_row - 1,
                           region.start_row + _MAX_TEMPLATE_SCAN_ROWS - 1)
            for start_row in range(region.start_row, last_row + 1):
                max_depth = min(_MAX_TEMPLATE_DEPTH,
                                region.end_row - start_row)
                for depth in range(1, max_depth + 1):
                    template = _template_at(
                        grid, region, start_row, depth, templates)
                    if template:
                        return (
                            HeaderCandidate(
                                start_row=start_row, depth=depth,
                                confidence=0.95, features={}, penalties=(),
                                source="template"),
                            template_columns(template),
                        )

        if decision.accepted:
            return decision.candidate, None
        return None, None

    def parse(self, path: Path) -> DocModel:
        document = Document(path)
        warnings: list[Warning] = []
        if document.element.body.findall(".//" + qn("w:ins")):
            warnings.append(Warning(
                code="revision_accepted", detail="已按接受修订处理"))
        blocks: list[Block] = []
        heading_path: list[str] = []
        table_no = 0
        llm_calls = 0

        for item in _iter_blocks(document):
            if isinstance(item, Paragraph):
                text = _paragraph_text(item)
                if not text:
                    continue
                level = _heading_level(item, text)
                if level is not None:
                    heading_path = heading_path[: level - 1] + [text]
                    blocks.append(Block(
                        block_id=f"b{len(blocks) + 1}", type="heading",
                        level=level, text=text, path=list(heading_path)))
                else:
                    blocks.append(Block(
                        block_id=f"b{len(blocks) + 1}", type="paragraph",
                        text=text, path=list(heading_path)))
                continue

            table_no += 1
            grid = _table_grid(item)
            if grid.nrows < 2 or grid.ncols < 2:
                continue
            region = Region(start_row=1, end_row=grid.nrows,
                            start_col=1, end_col=grid.ncols, density=1.0)
            decision = detect_header(grid, region, self.settings)
            header, forced_columns = self._resolve_header(grid, region, decision)

            if header is None and self.settings.header_llm_enabled \
                    and self.llm is not None \
                    and llm_calls < _MAX_LLM_HEADER_CALLS:
                header, raw = try_llm_header(grid, region, self.llm)
                llm_calls += 1
                if header is None:
                    warnings.append(Warning(
                        code="ambiguous_header", detail=raw[:500]))
            elif header is None and decision.candidate is not None:
                warnings.append(Warning(
                    code="ambiguous_header",
                    detail=f"置信度 {decision.candidate.confidence:.2f}"))

            table_key = f"t{table_no}"
            try:
                table, table_warnings = build_table(
                    grid, region, header, table_key=table_key,
                    settings=self.settings, forced_columns=forced_columns)
            except Exception as exc:  # 单表失败不拖垮整份文档
                warnings.append(Warning(
                    code="parse_partial", detail=f"table {table_key}: {exc}"))
                continue
            warnings.extend(
                Warning(code=w.code, detail=w.detail) for w in table_warnings)
            blocks.append(Block(
                block_id=f"b{len(blocks) + 1}", type="table", table=table,
                path=list(heading_path)))

        if not blocks:
            warnings.append(Warning(
                code="empty_document", detail="Word 中没有可入库内容"))

        title = (document.core_properties.title or "").strip() or path.stem
        return DocModel(
            parser=ParserInfo(name=self.name, version=self.version),
            meta=DocMeta(
                title=title,
                source_filename=path.name,
                sha256=sha256_file(path),
                effective_date=_header_footer_date(document),
            ),
            blocks=blocks,
            warnings=warnings,
        )


class DocParser:
    name = "doc"
    version = "0.1.0"

    def __init__(self, templates: list[dict] | None = None, llm=None,
                 settings: Settings | None = None):
        self._inner = DocxParser(
            templates=templates, llm=llm, settings=settings)
        self.settings = settings or get_settings()

    def supports(self, sniff) -> bool:
        return sniff.kind == "doc"

    def parse(self, path: Path) -> DocModel:
        try:
            converted = convert_document(path, "docx", timeout=120)
        except ConversionError as exc:
            raise ParseError(str(exc)) from exc
        model = self._inner.parse(converted)
        return model.model_copy(update={
            "parser": ParserInfo(name=self.name, version=self.version),
            "meta": model.meta.model_copy(update={
                "source_filename": path.name,
                "sha256": sha256_file(path),
            }),
        })
