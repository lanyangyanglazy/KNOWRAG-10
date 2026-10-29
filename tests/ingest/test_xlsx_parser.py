from types import SimpleNamespace

from ingest.config import Settings
from ingest.parsers.xlsx_parser import XlsxParser
from tests.ingest.builders import build_xlsx


class _CountingLLM:
    def __init__(self, content: str = "不是 JSON"):
        self.calls = 0
        self._content = content

    def invoke(self, messages):
        self.calls += 1
        return SimpleNamespace(content=self._content)


def _numeric_regions(regions: int) -> list[list[str]]:
    rows: list[list[str]] = []
    for i in range(regions):
        if rows:
            rows.append(["", ""])
        rows.append([str(i * 10 + 1), str(i * 10 + 2)])
        rows.append([str(i * 10 + 3), str(i * 10 + 4)])
    return rows


def test_single_table_auto_accepted(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsx", [
        ["序号", "部门", "金额"],
        [1, "综合处", 120],
        [2, "人事处", 98],
        [3, "财务处", 150],
    ], bold_rows=(1,))
    model = XlsxParser().parse(path)
    table = model.tables()[0]
    assert table.confidence >= 0.75
    assert table.source == "heuristic"
    assert [c.name for c in table.columns] == ["序号", "部门", "金额"]
    assert model.parser.name == "xlsx"


def test_multiple_regions_produce_two_tables(tmp_path):
    path = build_xlsx(tmp_path / "two.xlsx", [
        ["部门", "金额"],
        ["综合处", 120],
        ["", ""],
        ["姓名", "人数"],
        ["张伟", 3],
    ], bold_rows=(1, 4))
    model = XlsxParser().parse(path)
    assert len(model.tables()) == 2
    assert any(w.code == "multiple_tables" for w in model.warnings)


def test_no_header_is_flagged(tmp_path):
    path = build_xlsx(tmp_path / "raw.xlsx", [["1", "120"], ["2", "98"]])
    model = XlsxParser().parse(path)
    assert any(w.code == "table_no_header" for w in model.warnings)


def test_hidden_sheet_skipped(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "可见"
    wb.active.append(["部门", "金额"])
    wb.active.append(["综合处", 120])
    hidden = wb.create_sheet("隐藏")
    hidden.append(["部门", "金额"])
    hidden.append(["综合处", 1])
    hidden.sheet_state = "hidden"
    path = tmp_path / "hidden.xlsx"
    wb.save(path)
    model = XlsxParser().parse(path)
    assert len(model.tables()) == 1
    assert any(w.code == "hidden_sheet_skipped" for w in model.warnings)


def test_template_hit_sets_confidence(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsx", [
        ["部门", "Q1预算（万元）", "Q2预算（万元）"],
        ["综合处", 120, 130],
        ["人事处", 98, 101],
    ], bold_rows=(1,))
    template = {
        "id": "t1", "name": "预算模板", "scope_dept": None,
        "signature": {
            "ncols": 3, "header_depth": 1, "relative_header_row": 0,
            "columns_norm": ["部门", "q1预算", "q2预算"],
            "sheet_pattern": None, "sheet_name": "Sheet1",
        },
        "config": {"columns": [
            {"key": "c1", "name": "部门", "dtype": "string"},
            {"key": "c2", "name": "Q1预算", "dtype": "number", "unit": "万元"},
            {"key": "c3", "name": "Q2预算", "dtype": "number", "unit": "万元"},
        ]},
    }
    model = XlsxParser(templates=[template]).parse(path)
    table = model.tables()[0]
    assert table.source == "template"
    assert table.confidence == 0.95


def test_llm_header_calls_capped_per_workbook(tmp_path):
    # 裁决 2：每个工作簿最多 3 次 LLM 表头调用（只有 header_llm=always 才走）。
    path = build_xlsx(tmp_path / "numeric.xlsx", _numeric_regions(4))
    llm = _CountingLLM()
    settings = Settings(db_url="", upload_dir=tmp_path, header_llm="always")
    model = XlsxParser(llm=llm, settings=settings).parse(path)
    assert len(model.tables()) == 4
    assert llm.calls == 3
    assert sum(w.code == "ambiguous_header" for w in model.warnings) == 3


def test_llm_header_not_called_when_disabled(tmp_path):
    # 默认 HEADER_LLM=auto 在门禁前等同关闭：即使注入了 llm 也不得调用。
    path = build_xlsx(tmp_path / "numeric.xlsx", _numeric_regions(1))
    llm = _CountingLLM()
    settings = Settings(db_url="", upload_dir=tmp_path)
    XlsxParser(llm=llm, settings=settings).parse(path)
    assert llm.calls == 0


def test_template_hit_without_rule_candidates(tmp_path):
    # Minor 3：纯数字表头（2024|2025）不产生规则候选，模板仍须命中。
    path = build_xlsx(tmp_path / "years.xlsx", [
        ["2024", "2025"],
        [120, 130],
        [98, 101],
    ])
    template = {
        "id": "t-year", "name": "年度模板", "scope_dept": None,
        "signature": {
            "ncols": 2, "header_depth": 1, "relative_header_row": 0,
            "columns_norm": ["2024", "2025"],
            "sheet_pattern": None, "sheet_name": "Sheet1",
        },
        "config": {"columns": [
            {"key": "c1", "name": "2024年", "dtype": "number"},
            {"key": "c2", "name": "2025年", "dtype": "number"},
        ]},
    }
    model = XlsxParser(templates=[template]).parse(path)
    table = model.tables()[0]
    assert table.source == "template"
    assert table.confidence == 0.95
    assert table.header_rows == [1]
    assert [c.name for c in table.columns] == ["2024年", "2025年"]
    assert not any(w.code == "ambiguous_header" for w in model.warnings)


def test_template_width_mismatch_falls_back_to_rules(tmp_path):
    # Minor 3：模板 config 列数与区域不符时不得强塞给 build_table，
    # 否则表会因列数校验失败被 parse_partial 丢掉；应退回规则列。
    path = build_xlsx(tmp_path / "a.xlsx", [
        ["序号", "部门", "金额"],
        [1, "综合处", 120],
        [2, "人事处", 98],
        [3, "财务处", 150],
    ], bold_rows=(1,))
    template = {
        "id": "bad", "name": "坏模板", "scope_dept": None,
        "signature": {
            "ncols": 3, "header_depth": 1, "relative_header_row": 0,
            "columns_norm": ["序号", "部门", "金额"],
            "sheet_pattern": None, "sheet_name": "Sheet1",
        },
        "config": {"columns": [
            {"key": "c1", "name": "序号", "dtype": "number"},
            {"key": "c2", "name": "部门", "dtype": "string"},
        ]},
    }
    model = XlsxParser(templates=[template]).parse(path)
    assert len(model.tables()) == 1
    table = model.tables()[0]
    assert table.source == "heuristic"
    assert len(table.columns) == 3
    assert [c.name for c in table.columns] == ["序号", "部门", "金额"]
    assert not any(w.code == "parse_partial" for w in model.warnings)
