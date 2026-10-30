import hashlib
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ingest.config import Settings
from ingest.convert import ConversionError
from ingest.parsers.base import ParseError
from ingest.parsers.docx_parser import DocParser, DocxParser
from ingest.parsers.table_build import build_table as real_build_table
from tests.ingest.builders import build_docx


class _CountingLLM:
    def __init__(self, content: str = "不是 JSON"):
        self.calls = 0
        self._content = content

    def invoke(self, messages):
        self.calls += 1
        return SimpleNamespace(content=self._content)


def _budget_template(columns: list[dict] | None = None) -> dict:
    return {
        "id": "t1", "name": "预算模板", "scope_dept": None,
        "signature": {
            "ncols": 3, "header_depth": 1, "relative_header_row": 0,
            "columns_norm": ["序号", "部门", "金额"],
            "sheet_pattern": None, "sheet_name": None,
        },
        "config": {"columns": columns or [
            {"key": "c1", "name": "序号", "dtype": "number"},
            {"key": "c2", "name": "部门", "dtype": "string"},
            {"key": "c3", "name": "金额", "dtype": "number"},
        ]},
    }


def _three_col_table() -> list[list[str]]:
    return [
        ["序号", "部门", "金额"],
        ["1", "综合处", "120"],
        ["2", "人事处", "98"],
        ["3", "财务处", "150"],
    ]


def test_parses_headings_paragraphs_and_table(tmp_path):
    path = build_docx(
        tmp_path / "a.docx",
        paragraphs=[
            ("Heading 1", "二、预算管理"),
            (None, "各部门按季度编制预算。"),
            ("Heading 2", "（一）编制口径"),
            (None, "预算金额以万元为单位。"),
        ],
        table=[
            ["部门", "Q1预算（万元）", "Q2预算（万元）"],
            ["综合处", "120", "130"],
            ["人事处", "98", "101"],
        ],
    )
    model = DocxParser().parse(path)
    blocks = model.blocks
    assert [b.type for b in blocks] == ["heading", "paragraph", "heading", "paragraph", "table"]
    assert blocks[0].level == 1
    assert blocks[2].path == ["二、预算管理", "（一）编制口径"]
    table = model.tables()[0]
    assert [c.name for c in table.columns] == ["部门", "Q1预算", "Q2预算"]
    assert table.columns[1].unit == "万元"
    assert table.confidence >= 0.75
    assert blocks[4].path == ["二、预算管理", "（一）编制口径"]


def test_merged_cell_table(tmp_path):
    path = tmp_path / "merged.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    values = [["预算", "", ""], ["部门", "Q1", "Q2"], ["综合处", "120", "130"]]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            # (0,0)-(0,2) 合并后三格指向同一个 tc，给子格赋空串会清掉父格文本，
            # 因此合并区只在左上角写一次值。
            if r == 0 and c > 0:
                continue
            table.cell(r, c).text = value
    doc.save(path)

    model = DocxParser().parse(path)
    assert len(model.tables()) == 1
    names = [c.name for c in model.tables()[0].columns]
    assert names == ["预算/部门", "预算/Q1", "预算/Q2"]


def test_textbox_content_is_captured(tmp_path):
    path = tmp_path / "textbox.docx"
    doc = Document()
    doc.add_paragraph("正文段落")
    xml = (
        f'<w:p {nsdecls("w")}><w:r><w:t>另一段</w:t></w:r>'
        '<w:r><w:pict xmlns:v="urn:schemas-microsoft-com:vml"><v:shape>'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>文本框内容</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
    )
    doc.element.body.append(parse_xml(xml))
    doc.save(path)

    model = DocxParser().parse(path)
    texts = [b.text for b in model.blocks]
    assert any("文本框内容" in t for t in texts)


def test_empty_docx_is_flagged(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    model = DocxParser().parse(path)
    assert model.blocks == []
    assert any(w.code == "empty_document" for w in model.warnings)


def test_doc_conversion_failure_is_parse_error(tmp_path):
    path = tmp_path / "a.doc"
    path.write_bytes(b"fake")
    with patch("ingest.parsers.docx_parser.convert_document",
               side_effect=ConversionError("no soffice")):
        with pytest.raises(ParseError) as exc:
            DocParser().parse(path)
    assert "soffice" in str(exc.value)


def test_revision_is_warned(tmp_path):
    path = tmp_path / "rev.docx"
    doc = Document()
    doc.add_paragraph("正文")
    xml = (f'<w:p {nsdecls("w")}><w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
           '<w:r><w:t>插入</w:t></w:r></w:ins></w:p>')
    doc.element.body.append(parse_xml(xml))
    doc.save(path)
    model = DocxParser().parse(path)
    assert any(w.code == "revision_accepted" for w in model.warnings)


def test_revision_inserted_text_is_kept(tmp_path):
    # 设计 §7.2：接受修订后必须取插入文本；paragraph.text 不遍历 w:ins，
    # 只补警告不取字会让“已按接受修订处理”名不副实。
    path = tmp_path / "rev_insert.docx"
    doc = Document()
    doc.add_paragraph("正文")
    xml = (f'<w:p {nsdecls("w")}><w:r><w:t>前缀</w:t></w:r>'
           '<w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
           '<w:r><w:t>插入文本</w:t></w:r></w:ins></w:p>')
    doc.element.body.append(parse_xml(xml))
    doc.save(path)

    model = DocxParser().parse(path)
    assert any("插入文本" in b.text for b in model.blocks)
    assert any(w.code == "revision_accepted" for w in model.warnings)


def test_table_failure_is_isolated(tmp_path):
    # 裁决 4：单表构建失败只记 parse_partial，后续表格继续解析。
    path = tmp_path / "two_tables.docx"
    doc = Document()
    for _ in range(2):
        table = doc.add_table(rows=2, cols=2)
        for r, row in enumerate([["部门", "金额"], ["综合处", "120"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    doc.save(path)

    calls = {"count": 0}

    def flaky(sheet, region, header, *, table_key, settings, forced_columns=None):
        calls["count"] += 1
        if calls["count"] == 1:
            raise ValueError("boom")
        return real_build_table(sheet, region, header, table_key=table_key,
                                settings=settings, forced_columns=forced_columns)

    with patch("ingest.parsers.docx_parser.build_table", side_effect=flaky):
        model = DocxParser().parse(path)

    assert len(model.tables()) == 1
    partial = [w for w in model.warnings if w.code == "parse_partial"]
    assert len(partial) == 1
    assert "boom" in partial[0].detail


def test_table_truncation_keeps_columns_and_cuts_rows(tmp_path):
    # 上限截断只减行、不减列（列是表结构）；5 行 × 3 列、上限 9 → 前 3 行。
    path = build_docx(tmp_path / "tall.docx", paragraphs=[], table=[
        ["部门", "Q1预算（万元）", "Q2预算（万元）"],
        ["综合处", "120", "130"],
        ["人事处", "98", "101"],
        ["财务处", "150", "160"],
        ["办公室", "80", "90"],
    ])
    settings = Settings(db_url="", upload_dir=tmp_path, max_cells_per_sheet=9)
    model = DocxParser(settings=settings).parse(path)
    table = model.tables()[0]
    assert table.range == "A1:C3"
    assert len(table.columns) == 3
    assert any(w.code == "truncated" for w in model.warnings)


def test_template_beats_accepted_rules(tmp_path):
    # 裁决 1：规则本可接受该表，模板命中后 source/confidence 必须来自模板。
    path = build_docx(
        tmp_path / "rules.docx", paragraphs=[], table=_three_col_table())
    assert DocxParser().parse(path).tables()[0].source == "heuristic"
    table = DocxParser(
        templates=[_budget_template()]).parse(path).tables()[0]
    assert table.source == "template"
    assert table.confidence == 0.95


def test_template_width_mismatch_falls_back_to_heuristic(tmp_path):
    # 坏模板 config 列数与区域不符时不得塞给 build_table，应退回规则列。
    path = build_docx(
        tmp_path / "rules.docx", paragraphs=[], table=_three_col_table())
    template = _budget_template(columns=[
        {"key": "c1", "name": "序号", "dtype": "number"},
        {"key": "c2", "name": "部门", "dtype": "string"},
    ])
    model = DocxParser(templates=[template]).parse(path)
    table = model.tables()[0]
    assert table.source == "heuristic"
    assert len(table.columns) == 3
    assert not any(w.code == "parse_partial" for w in model.warnings)


def test_llm_header_calls_capped_per_document(tmp_path):
    # 裁决 2：纯数字表头不产生规则候选；每份文档最多 3 次 LLM 调用。
    path = tmp_path / "numeric_tables.docx"
    doc = Document()
    for _ in range(4):
        table = doc.add_table(rows=2, cols=2)
        for r, row in enumerate([["2024", "2025"], ["1", "2"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    doc.save(path)

    llm = _CountingLLM()
    settings = Settings(db_url="", upload_dir=tmp_path, header_llm="always")
    model = DocxParser(llm=llm, settings=settings).parse(path)
    assert len(model.tables()) == 4
    assert llm.calls == 3
    assert sum(w.code == "ambiguous_header" for w in model.warnings) == 3


def test_doc_success_path_uses_original_file_metadata(tmp_path):
    converted = tmp_path / "converted.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.sections[0].header.paragraphs[0].text = "2026-03-07"
    doc.save(converted)

    source = tmp_path / "legacy.doc"
    source.write_bytes(b"fake doc bytes")
    expected_sha = hashlib.sha256(source.read_bytes()).hexdigest()

    with patch("ingest.parsers.docx_parser.convert_document",
               return_value=converted):
        model = DocParser().parse(source)

    assert model.parser.name == "doc"
    assert model.meta.source_filename == "legacy.doc"
    assert model.meta.sha256 == expected_sha
    assert model.meta.effective_date == "2026-03-07"


def test_doc_converted_file_failure_is_parse_error(tmp_path):
    # 转换成功但产物损坏（python-docx 无法打开）也必须是 ParseError。
    source = tmp_path / "broken.doc"
    source.write_bytes(b"fake")
    converted = tmp_path / "broken.docx"
    converted.write_bytes(b"not a docx")
    with patch("ingest.parsers.docx_parser.convert_document",
               return_value=converted):
        with pytest.raises(ParseError):
            DocParser().parse(source)


def test_invalid_header_date_is_ignored(tmp_path):
    path = tmp_path / "bad_date.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.sections[0].header.paragraphs[0].text = "2026-13-45"
    doc.save(path)
    assert DocxParser().parse(path).meta.effective_date is None
