import zipfile
from dataclasses import replace
from pathlib import Path

import fitz
import pytest
from docx import Document

from ingest.config import get_settings
from ingest.errors import Quarantined, TooLarge, UnsupportedType
from ingest.sniff import sniff_file, strip_macros
from tests.ingest.builders import build_docx, build_xlsx

_VBA_REL = (
    b'<Relationship Id="rIdMacro" '
    b'Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
    b'Target="vbaProject.bin"/>'
)
_DOC_MAIN_CT = (b"application/vnd.openxmlformats-officedocument"
                b".wordprocessingml.document.main+xml")
_MACRO_DOC_MAIN_CT = b"application/vnd.ms-word.document.macroEnabled.main+xml"


def _build_macro_docx(path: Path, macro_main_type: bool = False) -> Path:
    """用 python-docx 造真实 DOCX，再注入 vbaProject 部件与关系（模拟 .docm 改名）。

    ``macro_main_type=True`` 时还把主文档部件的内容类型改写为 macroEnabled，
    模拟真正的 .docm 直接改名为 .docx 的情形（仅注入 vba 部件并不够）。
    """
    src = build_docx(path.with_suffix(".src.docx"), [(None, "hello")])
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(
            path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "[Content_Types].xml" and macro_main_type:
                assert _DOC_MAIN_CT in data
                data = data.replace(_DOC_MAIN_CT, _MACRO_DOC_MAIN_CT)
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>",
                                    _VBA_REL + b"</Relationships>")
            zout.writestr(info, data)
        zout.writestr("word/vbaProject.bin", b"macro-payload")
    src.unlink()
    return path


def test_sniffs_xlsx(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsx", [["部门", "金额"], ["综合处", 120]])
    result = sniff_file(path, "a.xlsx")
    assert result.kind == "xlsx"
    assert result.size_bytes > 0
    assert len(result.sha256) == 64
    assert result.macro is False


def test_extension_must_match_content(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsx", [["x"]])
    with pytest.raises(UnsupportedType):
        sniff_file(path, "a.docx")


def test_rejects_unknown_and_oversize(tmp_path):
    junk = tmp_path / "junk.bin"
    junk.write_bytes(b"not a document")
    with pytest.raises(UnsupportedType):
        sniff_file(junk, "junk.docx")

    path = build_xlsx(tmp_path / "a.xlsx", [["x"]])
    small = replace(get_settings(), max_file_mb=0)
    with pytest.raises(TooLarge):
        sniff_file(path, "a.xlsx", small)


def test_rejects_zip_bomb_ratio(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsx", [["x"]])
    with zipfile.ZipFile(path, "a") as zf:
        # 追加模式默认 ZIP_STORED，压缩比恒为 1；显式 DEFLATE 才是真实压缩炸弹
        zf.writestr("xl/media/big.bin", b"\0" * 30_000_000,
                    compress_type=zipfile.ZIP_DEFLATED)
    with pytest.raises(Quarantined):
        sniff_file(path, "a.xlsx")


def test_detects_and_strips_macros(tmp_path):
    path = build_xlsx(tmp_path / "a.xlsm", [["x"]])
    with zipfile.ZipFile(path, "a") as zf:
        zf.writestr("xl/vbaProject.bin", b"macro")
    assert sniff_file(path, "a.xlsm").macro is True

    clean = tmp_path / "clean.xlsm"
    assert strip_macros(path, clean) is True
    with zipfile.ZipFile(clean) as zf:
        assert "xl/vbaProject.bin" not in zf.namelist()


def test_strips_docx_macro_and_keeps_document_parseable(tmp_path):
    path = _build_macro_docx(tmp_path / "a.docx")
    result = sniff_file(path, "a.docx")
    assert result.kind == "docx"
    assert result.macro is True
    assert "macro_stripped" in result.warnings

    clean = tmp_path / "clean.docx"
    assert strip_macros(path, clean) is True
    with zipfile.ZipFile(clean) as zf:
        assert "word/vbaProject.bin" not in zf.namelist()
        rels = zf.read("word/_rels/document.xml.rels").lower()
    assert b"vbaproject" not in rels

    assert [p.text for p in Document(clean).paragraphs] == ["hello"]


def test_strips_renamed_docm_and_normalizes_content_type(tmp_path):
    path = _build_macro_docx(tmp_path / "disguised.docx", macro_main_type=True)
    assert sniff_file(path, "disguised.docx").macro is True

    clean = tmp_path / "clean.docx"
    assert strip_macros(path, clean) is True
    with zipfile.ZipFile(clean) as zf:
        content_types = zf.read("[Content_Types].xml")
    assert b"macroEnabled" not in content_types
    assert _DOC_MAIN_CT in content_types
    assert [p.text for p in Document(clean).paragraphs] == ["hello"]


def test_strip_macros_quarantines_broken_archive(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"PK\x03\x04not a real archive")
    out = tmp_path / "out.docx"
    with pytest.raises(Quarantined):
        strip_macros(broken, out)
    assert not out.exists()


def test_encrypted_pdf_is_quarantined(tmp_path):
    path = tmp_path / "secret.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "secret")
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    with pytest.raises(Quarantined):
        sniff_file(path, "secret.pdf")
