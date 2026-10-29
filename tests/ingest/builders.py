from pathlib import Path

from docx import Document
from openpyxl import Workbook


def build_xlsx(path: Path, rows: list[list], sheet: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def build_docx(path: Path, paragraphs: list[tuple[str | None, str]],
               table: list[list[str]] | None = None) -> Path:
    doc = Document()
    for style, text in paragraphs:
        doc.add_paragraph(text, style=style)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.save(path)
    return path
