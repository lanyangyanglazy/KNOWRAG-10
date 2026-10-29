import zipfile
from dataclasses import replace
from datetime import datetime

from openpyxl import Workbook

from ingest.config import get_settings
from ingest.parsers.xlsx_reader import has_pivot_tables, read_workbook
from tests.ingest.builders import build_xlsx


def test_reads_display_values_styles_and_flags(tmp_path):
    path = tmp_path / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "预算"
    ws.append(["部门", "占比", "日期", "金额"])
    ws.append(["综合处", 0.15, datetime(2026, 1, 1), 1234.5])
    ws["B2"].number_format = "0%"
    ws["C2"].number_format = "yyyy-mm-dd"
    ws["D2"].number_format = "#,##0.00"
    ws["A1"].font = ws["A1"].font.copy(bold=True)
    ws["A2"].fill = ws["A2"].fill.copy(fill_type="solid")
    ws["E5"] = "=B2+1"
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["B"].hidden = True
    wb.save(path)

    grid = read_workbook(path)
    sheet = grid.sheets[0]
    assert sheet.name == "预算"
    assert sheet.cell(2, 2).text == "15%"
    assert sheet.cell(2, 3).text == "2026-01-01"
    assert sheet.cell(2, 4).text == "1,234.50"
    assert sheet.cell(1, 1).bold is True
    assert sheet.cell(2, 1).has_fill is True
    assert sheet.cell(2, 2).is_numeric is True
    assert sheet.cell(2, 1).hidden is True
    assert sheet.cell(5, 5).formula == "=B2+1"
    assert any(w["code"] == "formula_not_cached" for w in grid.warnings)
    assert any(w["code"] == "hidden_rows" for w in grid.warnings)


def test_expands_merged_cells(tmp_path):
    path = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "预算"
    ws.merge_cells("A1:B1")
    ws["A2"] = "部门"
    ws["B2"] = "Q1"
    ws["A3"] = "综合处"
    ws["B3"] = 120
    wb.save(path)

    sheet = read_workbook(path).sheets[0]
    assert sheet.merges == ["A1:B1"]
    assert sheet.cell(1, 2).text == "预算"


def test_hidden_sheet_is_kept_with_flag(tmp_path):
    path = tmp_path / "hidden.xlsx"
    wb = Workbook()
    wb.active.title = "可见"
    wb.active["A1"] = "x"
    hidden = wb.create_sheet("隐藏")
    hidden["A1"] = "y"
    hidden.sheet_state = "hidden"
    wb.save(path)

    sheets = read_workbook(path).sheets
    assert [(s.name, s.hidden) for s in sheets] == [("可见", False), ("隐藏", True)]


def test_truncates_at_cell_limit(tmp_path):
    path = tmp_path / "big.xlsx"
    wb = Workbook()
    for r in range(1, 6):
        for c in range(1, 6):
            wb.active.cell(r, c, f"{r}-{c}")
    wb.save(path)

    settings = replace(get_settings(), max_cells_per_sheet=10)
    grid = read_workbook(path, settings)
    assert grid.sheets[0].nrows * grid.sheets[0].ncols <= 10
    assert any(w["code"] == "truncated" for w in grid.warnings)


def test_truncation_keeps_columns_and_cuts_rows(tmp_path):
    # 50000 行 × 10 列这类“高表”：优先截行保列，丢列会破坏 Table/切片语义
    path = build_xlsx(tmp_path / "tall.xlsx",
                      [[f"{r}-{c}" for c in range(1, 11)] for r in range(1, 301)])

    settings = replace(get_settings(), max_cells_per_sheet=1000)
    grid = read_workbook(path, settings)

    sheet = grid.sheets[0]
    assert sheet.ncols == 10
    assert sheet.nrows == 100
    assert sheet.cell(100, 10).text == "100-10"
    assert any(w["code"] == "truncated" for w in grid.warnings)


def test_truncation_cuts_columns_only_when_one_row_exceeds_budget(tmp_path):
    # 单行就超预算（10 列 > 8 格）时才允许截列
    path = build_xlsx(tmp_path / "wide.xlsx",
                      [[f"{c}" for c in range(1, 11)]])

    settings = replace(get_settings(), max_cells_per_sheet=8)
    grid = read_workbook(path, settings)

    sheet = grid.sheets[0]
    assert sheet.nrows == 1
    assert sheet.ncols == 8
    assert any(w["code"] == "truncated" for w in grid.warnings)


def test_detects_pivot_tables(tmp_path):
    pivot = tmp_path / "pivot.xlsx"
    with zipfile.ZipFile(pivot, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/pivotTables/pivotTable1.xml", "<pivotTable/>")
    assert has_pivot_tables(pivot) is True

    plain = tmp_path / "plain.xlsx"
    Workbook().save(plain)
    assert has_pivot_tables(plain) is False
